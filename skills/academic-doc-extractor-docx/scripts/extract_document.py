#!/usr/bin/env python3
"""
Document Data Extractor for Academic Literature Review
Extracts structured text, metadata, tables, and mathematical formulas from PDF and Word (.docx) files.
"""

import sys
import os
import json
import re
from typing import Dict, List, Any, Optional

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None


def extract_from_pdf(file_path: str) -> Dict[str, Any]:
    if not fitz:
        raise ImportError("PyMuPDF (fitz) is not installed. Please install it with 'pip install pymupdf'.")
    
    doc = fitz.open(file_path)
    metadata = {
        "title": doc.metadata.get("title") or os.path.basename(file_path),
        "author": doc.metadata.get("author") or "Unknown",
        "subject": doc.metadata.get("subject") or "",
        "keywords": doc.metadata.get("keywords") or "",
        "page_count": len(doc),
        "file_path": os.path.abspath(file_path),
        "file_type": "PDF"
    }

    pages_content = []
    full_text_list = []
    tables = []
    
    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text("text")
        full_text_list.append(text)
        
        # Extract tables if available in PyMuPDF
        try:
            tabs = page.find_tables()
            if tabs.tables:
                for idx, t in enumerate(tabs.tables):
                    df_data = t.extract()
                    if df_data:
                        tables.append({
                            "page": page_num + 1,
                            "table_index": idx + 1,
                            "headers": df_data[0] if len(df_data) > 0 else [],
                            "rows": df_data[1:] if len(df_data) > 1 else []
                        })
        except Exception:
            pass
            
        pages_content.append({
            "page_number": page_num + 1,
            "char_count": len(text),
            "text": text
        })

    doc.close()
    combined_text = "\n\n".join(full_text_list)
    
    # Identify structured sections and candidate formulas
    sections = parse_sections_from_text(combined_text)
    formulas = detect_formulas(combined_text)

    return {
        "metadata": metadata,
        "sections": sections,
        "tables": tables,
        "formulas": formulas,
        "full_text": combined_text
    }


def extract_from_docx(file_path: str) -> Dict[str, Any]:
    if not docx:
        raise ImportError("python-docx is not installed. Please install it with 'pip install python-docx'.")
    
    doc = docx.Document(file_path)
    core_props = doc.core_properties
    metadata = {
        "title": core_props.title or os.path.basename(file_path),
        "author": core_props.author or "Unknown",
        "subject": core_props.subject or "",
        "keywords": core_props.keywords or "",
        "file_path": os.path.abspath(file_path),
        "file_type": "DOCX",
        "paragraphs_count": len(doc.paragraphs),
        "tables_count": len(doc.tables)
    }

    sections = []
    current_section = {"title": "Introduction / Preamble", "content": []}
    full_text_list = []

    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        full_text_list.append(txt)
        
        # Check if paragraph has a heading style
        if p.style and p.style.name.startswith("Heading"):
            if current_section["content"]:
                current_section["content"] = "\n".join(current_section["content"])
                sections.append(current_section)
            current_section = {"title": txt, "style": p.style.name, "content": []}
        else:
            current_section["content"].append(txt)

    if current_section["content"]:
        if isinstance(current_section["content"], list):
            current_section["content"] = "\n".join(current_section["content"])
        sections.append(current_section)

    tables = []
    for t_idx, table in enumerate(doc.tables):
        table_rows = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            table_rows.append(row_data)
        if table_rows:
            tables.append({
                "table_index": t_idx + 1,
                "headers": table_rows[0],
                "rows": table_rows[1:]
            })

    combined_text = "\n\n".join(full_text_list)
    formulas = detect_formulas(combined_text)

    return {
        "metadata": metadata,
        "sections": sections,
        "tables": tables,
        "formulas": formulas,
        "full_text": combined_text
    }


def parse_sections_from_text(text: str) -> List[Dict[str, str]]:
    """Heuristic identification of common academic headings."""
    heading_regex = re.compile(
        r'^(?:(?:\d+\.|\d+\.\d+|\b[IVXLCDM]+\.|\bSection\s+\d+)\s+)?'
        r'(Abstract|Introduction|Background|Related Work|Literature Review|Methodology|Methods|'
        r'Theory|Model|Experimental Setup|Results|Discussion|Conclusion|References)\b.*$',
        re.IGNORECASE | re.MULTILINE
    )
    
    matches = list(heading_regex.finditer(text))
    sections = []
    if not matches:
        return [{"title": "Full Content", "content": text}]
    
    for i in range(len(matches)):
        start = matches[i].start()
        end = matches[i+1].start() if i + 1 < len(matches) else len(text)
        title = matches[i].group(0).strip()
        body = text[matches[i].end():end].strip()
        sections.append({"title": title, "content": body})
        
    return sections


def detect_formulas(text: str) -> List[str]:
    """Detect potential LaTeX or mathematical formula patterns."""
    formula_patterns = [
        r'\$\$([\s\S]+?)\$\$',       # Display math
        r'\$([^\$\n]+?)\$',           # Inline math
        r'\\begin\{(?:equation|align|gather)\*?\}[\s\S]+?\\end\{(?:equation|align|gather)\*?\}',
        r'(?:[A-Za-z0-9_]+(?:\^[A-Za-z0-9_]+)?\s*=\s*[^.\n]{4,}\b)' # Basic algebraic equality
    ]
    formulas = []
    for pattern in formula_patterns:
        for match in re.finditer(pattern, text):
            f_str = match.group(0).strip()
            if len(f_str) > 4 and f_str not in formulas:
                formulas.append(f_str)
    return formulas[:50]  # Cap to top 50 to prevent bloating


def main():
    if len(sys.argv) < 2:
        print("Usage: python extract_document.py <path_to_pdf_or_docx> [--json] [--output <out_path>]")
        sys.exit(1)
        
    input_path = sys.argv[1]
    if not os.path.exists(input_path):
        print(f"Error: File not found: {input_path}")
        sys.exit(1)
        
    ext = os.path.splitext(input_path)[1].lower()
    if ext == ".pdf":
        result = extract_from_pdf(input_path)
    elif ext in [".docx", ".doc"]:
        result = extract_from_docx(input_path)
    else:
        print(f"Error: Unsupported file format '{ext}'. Must be .pdf or .docx.")
        sys.exit(1)
        
    output_path = None
    if "--output" in sys.argv:
        idx = sys.argv.index("--output")
        if idx + 1 < len(sys.argv):
            output_path = sys.argv[idx + 1]

    if "--json" in sys.argv or (output_path and output_path.endswith(".json")):
        json_str = json.dumps(result, indent=2, ensure_ascii=False)
        if output_path:
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(json_str)
            print(f"Extracted JSON saved to: {output_path}")
        else:
            print(json_str)
    else:
        # Formatted readable summary
        meta = result["metadata"]
        summary_lines = [
            f"# Document Extraction Summary: {meta.get('title')}",
            f"- **File Type:** {meta.get('file_type')}",
            f"- **Author:** {meta.get('author')}",
            f"- **Sections Extracted:** {len(result.get('sections', []))}",
            f"- **Tables Extracted:** {len(result.get('tables', []))}",
            f"- **Candidate Formulas Detected:** {len(result.get('formulas', []))}\n",
            "## Section Breakdown:"
        ]
        for s in result.get("sections", []):
            summary_lines.append(f"### {s['title']}")
            summary_lines.append(s['content'][:300] + ("..." if len(s['content']) > 300 else "") + "\n")
            
        full_summary = "\n".join(summary_lines)
        if output_path:
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(full_summary)
            print(f"Summary saved to: {output_path}")
        else:
            print(full_summary)


if __name__ == "__main__":
    main()
